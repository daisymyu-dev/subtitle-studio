"""Offline bilingual DOCX export and non-overwriting OS Desktop copies."""

import ctypes
from ctypes import wintypes
from copy import deepcopy
import io
import math
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from .subtitles import safe_filename


_INVALID_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def _text(value, field):
    if not isinstance(value, str) or not value.strip() or _INVALID_XML.search(value):
        raise ValueError(f'{field} must be nonempty XML-compatible text.')
    return value


def _payload(title, url, blocks, source_note, english_note):
    data = dict(title=_text(title, 'title'), url=_text(url, 'url'),
                source_note=_text(source_note, 'source_note'),
                english_note=_text(english_note, 'english_note'), blocks=[])
    if not isinstance(blocks, list) or not blocks:
        raise ValueError('blocks must be a nonempty list of complete bilingual blocks.')
    for index, block in enumerate(blocks):
        if not isinstance(block, dict):
            raise ValueError(f'Block {index + 1} must contain start, end, english and chinese.')
        start, end = block.get('start'), block.get('end')
        try:
            invalid_time = (any(isinstance(value, bool) or not isinstance(value, (int, float))
                                or not math.isfinite(value * 1000) for value in (start, end))
                            or start < 0 or end <= start)
        except OverflowError:
            invalid_time = True
        if invalid_time:
            raise ValueError(f'Block {index + 1} must have a finite nonnegative time range.')
        data['blocks'].append(dict(start=start, end=end,
                                   english=_text(block.get('english'), f'Block {index + 1} english'),
                                   chinese=_text(block.get('chinese'), f'Block {index + 1} chinese'),
                                   fragment=block.get('fragment') is True,
                                   fragment_note='结尾完整性待核对' if 'restored_tail_uncertain' in str(block.get('fragment_reason','')) else '残句 / 截断片段'))
    return data


def _timestamp(seconds):
    # Match the former builder's nonnegative, half-up millisecond rounding.
    milliseconds = math.floor(seconds * 1000 + .5)
    hours, milliseconds = divmod(milliseconds, 3600000)
    minutes, milliseconds = divmod(milliseconds, 60000)
    seconds, milliseconds = divmod(milliseconds, 1000)
    return f'{hours:02}:{minutes:02}:{seconds:02}.{milliseconds:03}'


def _font(style, size):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.get_or_add_rPr()
    # Built-in styles carry theme fonts, which otherwise override explicit names.
    properties.rFonts.attrib.clear()
    for name, value in (('ascii', 'Arial'), ('hAnsi', 'Arial'),
                        ('eastAsia', 'Microsoft YaHei'), ('cs', 'Arial')):
        properties.rFonts.set(qn('w:' + name), value)
    complex_size = properties.find(qn('w:szCs'))
    if complex_size is None:
        complex_size = OxmlElement('w:szCs')
        properties.insert(properties.index(properties.sz) + 1, complex_size)
    complex_size.set(qn('w:val'), str(size * 2))
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'en-US')
    language.set(qn('w:eastAsia'), 'zh-CN')
    properties.append(language)


def _document_bytes(data):
    """Build locally with the project dependency; no tools or profile discovery."""
    document = Document()
    styles = document.styles
    normal = styles['Normal']
    _font(normal, 11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    # Readers not applying the Normal style still get the same readable default.
    default_run = styles.element.xpath('w:docDefaults/w:rPrDefault')[0]
    default_run.replace(default_run[0], deepcopy(normal.element.rPr))

    title_style = styles['Title']
    for border in title_style.element.xpath('w:pPr/w:pBdr'):
        border.getparent().remove(border)
    _font(title_style, 20)
    title_style.font.bold = True
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.keep_with_next = True

    for name, size, before, after, keep_next in (
            ('Provenance', 10, 0, 5, False),
            ('TimeRange', 10, 10, 3, True),
            ('English', 11, 0, 3, True),
            ('Chinese', 11, 0, 7, False)):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        _font(style, size)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = keep_next
    for name, next_name in (('Provenance', 'Normal'), ('TimeRange', 'English'),
                            ('English', 'Chinese'), ('Chinese', 'TimeRange')):
        styles[name].next_paragraph_style = styles[next_name]
    styles['TimeRange'].font.bold = True
    styles['TimeRange'].paragraph_format.keep_together = True

    section = document.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.top_margin = section.bottom_margin = Twips(1134)
    section.left_margin = section.right_margin = Twips(1134)
    section.footer_distance = Twips(567)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Page ')
    for field_name, following in (('PAGE', ' of '), ('NUMPAGES', '')):
        run = footer.add_run()
        for tag, value in (('w:fldChar', 'begin'), ('w:instrText', field_name),
                           ('w:fldChar', 'separate'), ('w:t', '1'), ('w:fldChar', 'end')):
            element = OxmlElement(tag)
            if tag == 'w:fldChar':
                element.set(qn('w:fldCharType'), value)
            else:
                element.text = value
            run.element.append(element)
        footer.add_run(following)
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0, 0, 0)

    def paragraph(text, style):
        # python-docx creates real w:br/w:tab; normalize CRLF to one line break.
        document.add_paragraph(text.replace('\r\n', '\n').replace('\r', '\n'), style)

    paragraph(data['title'], 'Title')
    paragraph(f"来源链接：{data['url']}", 'Provenance')
    paragraph(f"英文来源：{data['english_note']}", 'Provenance')
    paragraph(f"中文来源：{data['source_note']}", 'Provenance')
    for block in data['blocks']:
        time_range = f"{_timestamp(block['start'])} --> {_timestamp(block['end'])}"
        if block['fragment']:
            time_range += '  · ' + block['fragment_note']
        paragraph(time_range, 'TimeRange')
        paragraph(block['english'], 'English')
        paragraph(block['chinese'], 'Chinese')

    document.core_properties.author = 'Subtitle Studio'
    # The high-level metadata setter imposes a 255-character limit absent from
    # the former builder/API. This is already validated text, not XML markup.
    document.core_properties._element.get_or_add_title().text = data['title']
    document.core_properties.comments = (
        'English transcript paired with Chinese text and source provenance')
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def write_word_document(destination, title, url, blocks, source_note, english_note) -> Path:
    """Write every block to the exact .docx path and return that Path.

    Text is serialized in-process with python-docx, never shell code or XML markup.
    A completed DOCX is staged beside the destination then atomically replaced;
    build/write failures leave an existing destination untouched. No translation
    or network access is performed, and the supplied title is never a filename.
    """
    destination = Path(destination)
    if destination.suffix.lower() != '.docx':
        raise ValueError('destination must be an exact .docx file path.')
    data = _payload(title, url, blocks, source_note, english_note)
    try:
        content = _document_bytes(data)
    except Exception:
        # Serializer errors can contain caption text; do not expose their chain.
        raise RuntimeError('Local Word document generation failed.') from None
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if (not {'[Content_Types].xml', 'word/document.xml'}.issubset(archive.namelist())
                    or archive.testzip() is not None):
                raise BadZipFile
    except BadZipFile:
        raise RuntimeError('Local Word builder did not return a complete DOCX.') from None

    destination.parent.mkdir(parents=True, exist_ok=True)
    staged = None
    try:
        with tempfile.NamedTemporaryFile(mode='wb', prefix='.word-', suffix='.tmp',
                                         dir=destination.parent, delete=False) as output:
            staged = Path(output.name)
            output.write(content)
        os.replace(staged, destination)
    finally:
        if staged is not None:
            staged.unlink(missing_ok=True)
    return destination


def _desktop_directory():
    if sys.platform != 'win32':
        return Path.home() / 'Desktop'
    get_folder = ctypes.windll.shell32.SHGetFolderPathW
    get_folder.argtypes = [wintypes.HWND, ctypes.c_int, wintypes.HANDLE,
                           wintypes.DWORD, wintypes.LPWSTR]
    get_folder.restype = ctypes.c_long
    buffer = ctypes.create_unicode_buffer(32768)
    # CSIDL_DESKTOPDIRECTORY respects localized, redirected and OneDrive Desktops.
    status = get_folder(None, 0x0010, None, 0, buffer)
    if status != 0 or not buffer.value:
        raise OSError('Windows could not locate the current user Desktop directory.')
    return Path(buffer.value)


def save_to_desktop(path, desktop_dir=None) -> Path:
    """Copy to the OS Desktop with exclusive creation and numbered collisions.

    desktop_dir is an optional trusted caller/test override. Existing files are
    never opened for writing; any newly created partial copy is removed on error.
    The source is never moved, deleted or overwritten, including on the Desktop.
    """
    source = Path(path)
    with source.open('rb') as reader:
        desktop = Path(desktop_dir) if desktop_dir is not None else _desktop_directory()
        desktop.mkdir(parents=True, exist_ok=True)
        stem = safe_filename(source.stem)
        number = 0
        while True:
            suffix = f' ({number})' if number else ''
            candidate = desktop / f'{stem}{suffix}{source.suffix}'
            try:
                output = candidate.open('xb')
            except FileExistsError:
                number += 1
                continue
            try:
                with output:
                    shutil.copyfileobj(reader, output)
            except BaseException:
                candidate.unlink(missing_ok=True)
                raise
            return candidate
