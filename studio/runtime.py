"""Separate read-only program assets from writable local application state."""

import os
from pathlib import Path
import sys
import tempfile

VERSION = '0.1.0-beta.1'


def asset_root():
    """PyInstaller's onedir asset directory; never the current directory."""
    if getattr(sys, 'frozen', False):
        return Path(sys._MEIPASS).resolve()
    return Path(__file__).resolve().parents[1]


def data_root(override=None):
    """CLI > environment > per-user frozen data / historical source checkout."""
    value = override if override is not None else os.environ.get('SUBTITLE_STUDIO_DATA_DIR')
    frozen = getattr(sys, 'frozen', False)
    if value is None:
        if not frozen:
            return asset_root()
        local = os.environ.get('LOCALAPPDATA')
        root = Path(local) if local else Path.home() / 'AppData' / 'Local'
        value = root / 'SubtitleStudio'
    if not str(value).strip():
        raise ValueError('The data directory must not be empty.')
    root = Path(value).expanduser().resolve()
    if root == Path(root.anchor) or str(root).startswith('\\\\'):
        raise ValueError('Choose a local data folder, not a drive root or network share.')
    if frozen:
        program = Path(sys.executable).resolve().parent
        if root.is_relative_to(program) or root.is_relative_to(asset_root()):
            raise ValueError('Data must be outside the installed program directory.')
    if root.exists() and not root.is_dir():
        raise ValueError('The data directory points to a file.')
    return root


def prepare_data(root):
    """Keep temporary files and future model downloads under the selected root."""
    root = Path(root)
    temporary = root / '.work' / 'tmp'
    temporary.mkdir(parents=True, exist_ok=True)
    os.environ['TEMP'] = os.environ['TMP'] = str(temporary)
    tempfile.tempdir = str(temporary)
    # Set these before importing the server/optional Hugging Face libraries.
    os.environ['HF_HOME'] = str(root / '.models' / 'huggingface')
    os.environ['HF_HUB_CACHE'] = str(root / '.models' / 'huggingface' / 'hub')
    os.environ['HUGGINGFACE_HUB_CACHE'] = os.environ['HF_HUB_CACHE']
    os.environ['XDG_CACHE_HOME'] = str(root / '.cache')


def check_runtime():
    """Exercise packaged libraries locally, without model weights or downloads."""
    import importlib.util
    import io
    import subprocess
    import wave
    from docx import Document
    from yt_dlp.version import __version__ as yt_version
    import huggingface_hub
    import onnxruntime
    import tokenizers

    output = io.BytesIO()
    document = Document()
    document.add_paragraph('Subtitle Studio runtime check.')
    document.save(output)
    if Document(io.BytesIO(output.getvalue())).paragraphs[0].text != 'Subtitle Studio runtime check.':
        raise RuntimeError('DOCX runtime verification failed.')
    # Exercise the tokenizer's native extension without external model assets.
    tokenizer = tokenizers.Tokenizer(tokenizers.models.WordLevel({'hello': 0, '[UNK]': 1}, unk_token='[UNK]'))
    if tokenizer.encode('hello').ids != [0] or 'CPUExecutionProvider' not in onnxruntime.get_available_providers():
        raise RuntimeError('Punctuation native runtime verification failed.')
    report = {'version': VERSION, 'docx': True, 'yt_dlp': yt_version, 'asr': None, 'node': None,
              'core_ready': True, 'punctuation': {'onnxruntime': onnxruntime.__version__,
                  'tokenizers': tokenizers.__version__, 'huggingface_hub': huggingface_hub.__version__}}
    if importlib.util.find_spec('faster_whisper'):
        import av
        import ctranslate2
        from faster_whisper.audio import decode_audio
        audio = io.BytesIO()
        with wave.open(audio, 'wb') as stream:
            stream.setnchannels(1)
            stream.setsampwidth(2)
            stream.setframerate(16000)
            stream.writeframes(b'\0\0' * 16000)
        audio.seek(0)
        report['asr'] = {'decoded_samples': len(decode_audio(audio)), 'av': av.__version__,
                         'ctranslate2': ctranslate2.__version__, 'onnxruntime': onnxruntime.__version__,
                         'tokenizers': tokenizers.__version__}
    node = asset_root() / 'runtime' / 'node.exe'
    if getattr(sys, 'frozen', False) and node.is_file():
        result = subprocess.run([str(node), '--version'], capture_output=True, text=True,
                                check=True, timeout=15, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        report['node'] = result.stdout.strip()
    return report
